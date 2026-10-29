"""One editable Word master per artifact; PDF is exported from those masters."""
from pathlib import Path
import json,re,sys,copy,os
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
H=Path(__file__).resolve().parent; OUT=H.parent;OUT.mkdir(exist_ok=True)

from latex2mathml.converter import convert
xslt=etree.XSLT(etree.parse(str(Path(os.environ.get('PROGRAMFILES','C:/Program Files'))/'Microsoft Office/root/Office16/MML2OMML.XSL')))
tables=json.loads((H/'original_tables.json').read_text('utf8'))
refs=json.loads((H/'venue_reference_audit.json').read_text('utf8'))['references']
def author_key(r):
    first=r['fields']['author'].split('; ')[0].split(' and ')[0]
    if ', ' in first:
        surname,given=first.split(', ',1)
    else:
        parts=first.split();surname,given=parts[-1],' '.join(parts[:-1])
    return surname.casefold(),given.casefold(),r['fields']['year'],r['fields']['title'].casefold()
refs.sort(key=author_key)
ref_numbers={r['number']:i+1 for i,r in enumerate(refs)}
def cite(text):
    return re.sub(r'\[(\d+(?:,\s*\d+)*)\]',lambda m:'['+', '.join(str(n) for n in sorted(ref_numbers[int(v)] for v in m[1].split(',')))+']',text)
FULL=6.875; COL=3.28125

def el(name,**kwargs):
    x=OxmlElement(name)
    for k,v in kwargs.items():x.set(qn('w:'+k),str(v))
    return x

def setup(doc,supp=False):
    n=doc.styles['Normal'];n.font.name='Times New Roman';n.font.size=Pt(10)
    n.paragraph_format.line_spacing=Pt(12);n.paragraph_format.space_after=Pt(0)
    n.paragraph_format.widow_control=True;n.paragraph_format.first_line_indent=Inches(.13)
    n.paragraph_format.keep_together=False;n.paragraph_format.keep_with_next=False
    ln=doc.styles['Line Number'] if 'Line Number' in doc.styles else doc.styles.add_style('Line Number',WD_STYLE_TYPE.CHARACTER,builtin=True)
    ln.font.name='Times New Roman';ln.font.size=Pt(7);ln.font.color.rgb=RGBColor(90,90,90)
    for name,size in [('Heading 1',12),('Heading 2',10),('Title',14),('Caption',9)]:
        if name not in doc.styles:continue
        st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.first_line_indent=Inches(0);st.paragraph_format.space_after=Pt(4);st.paragraph_format.space_before=Pt(8)
        if name in ['Heading 1','Heading 2']:st.font.bold=True;st.paragraph_format.keep_with_next=True
    doc.styles['Title'].paragraph_format.space_after=Pt(10)
    doc.styles['Title'].paragraph_format.space_before=Pt(0)
    for st in doc.styles:
        for b in list(st.element.findall('.//'+qn('w:pBdr'))):b.getparent().remove(b)
        for f in st.element.findall('.//'+qn('w:rFonts')):
            for a in list(f.attrib):
                if 'Theme' in a or 'theme' in a:del f.attrib[a]
    doc.styles['Caption'].paragraph_format.space_before=Pt(2);doc.styles['Caption'].paragraph_format.space_after=Pt(6)
    doc.styles['Caption'].font.bold=False
    doc.styles['Caption'].paragraph_format.line_spacing=Pt(10.5)
    sec=doc.sections[0];configure_section(sec,1)
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Inches(0)
    field=el('w:fldSimple',instr='PAGE');p._p.append(field)
    doc.core_properties.author='Anonymous';doc.core_properties.last_modified_by='Anonymous';doc.core_properties.comments=''
    doc.core_properties.subject='Camera-space human mesh refinement'
    doc.settings.element.append(el('w:autoHyphenation',val='true'))
    doc.settings.element.append(el('w:hyphenationZone',val='180'))
    doc.settings.element.append(el('w:consecutiveHyphenLimit',val='2'))

def configure_section(sec,cols):
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(1);sec.bottom_margin=Inches(1.125);sec.left_margin=sec.right_margin=Inches(.8125)
    sec.header_distance=Inches(.4);sec.footer_distance=Inches(.55)
    c=sec._sectPr.find(qn('w:cols'));c.set(qn('w:num'),str(cols));c.set(qn('w:space'),'450')
    ln=sec._sectPr.find(qn('w:lnNumType'))
    if ln is None:ln=el('w:lnNumType',countBy='1',start='1',distance='100',restart='newPage');sec._sectPr.append(ln)

def section(doc,cols,new=False):
    sec=doc.add_section(WD_SECTION.NEW_PAGE if new else WD_SECTION.CONTINUOUS);configure_section(sec,cols)
    p=doc.paragraphs[-1];p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(1)
    for r in p.runs:r.font.size=Pt(1)

def para(doc,text,style=None,boldlead=False):
    text=cite(text)
    p=doc.add_paragraph(style=style);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if boldlead and '. ' in text:
        a,b=text.split('. ',1);p.add_run(a+'. ').bold=True;p.add_run(b)
    else:p.add_run(text)
    return p

def eq(doc,tex):
    tex=tex.replace('\\\\','\\')
    if len(tex)>105 and '\\quad' in tex:tex='\\begin{aligned}'+tex.replace(',\\quad',',\\\\').replace('\\quad','\\\\')+'\\end{aligned}'
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0)
    p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1
    p.paragraph_format.keep_together=True
    om=xslt(etree.fromstring(convert(tex).encode('utf8'))).getroot()
    for r in om.findall('.//'+qn('m:r')):
        pr=OxmlElement('w:rPr');pr.append(el('w:rFonts',ascii='Cambria Math',hAnsi='Cambria Math'));pr.append(el('w:sz',val='18'));r.insert(0,pr)
    p._p.append(om)

def figure(doc,b,num,supp=False,open_wide=True,close_wide=True):
    key=b['key'];wide=key in ['overview','identity','sequences','dynamics'] and not supp
    if wide and open_wide:section(doc,1)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0)
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=True
    p.paragraph_format.line_spacing=1
    filenames={'overview':'overview_final.png','identity':'identity_final.png','motion':'motion_gemini_style_final.png'}
    f=H/'figures'/filenames.get(key,key+'.png')
    if key=='identity' and not f.exists():f=H/'figures'/'identity_gemini_style_v2.png'
    width=6.0 if supp and key=='motion' else (FULL if wide or supp else COL)
    p.add_run().add_picture(str(f),width=Inches(width))
    cap=para(doc,f'Figure {"S" if supp else ""}{num}. '+b['caption'],'Caption');cap.paragraph_format.keep_together=True
    if wide and close_wide:section(doc,2)

def format_num(v):
    try:
        f=float(v)
        if '.' not in v:return v
        return f'{f:.4f}' if 0<abs(f)<1 else f'{f:.2f}'
    except:return v

def table(doc,b,num,supp=False):
    key=b['key']
    if key=='published_benchmark':
        bm=json.loads((H/'published_benchmark_rows.json').read_text('utf8'))
        data=[bm['columns']]+[r['display'] for r in bm['rows']]
    else:
        idx={'interfaces':0,'loss':1,'main':2,'sequences':3,'ema':4}[key]
        data=copy.deepcopy(tables[idx])
    if key=='main':
        tm=json.loads((H/'table_manifest.json').read_text('utf8'))
        data[0]=['Model / refinement stage','MPJPE ↓','PA-MPJPE ↓','PVE ↓','Accel. ↓']
        data[1:]=[[r[0]+'\n'+r[1]]+[f'{float(v):.2f}' for v in r[2:]] for r in tables[idx][1:]]
        for row,source in [(2,'parent'),(3,'mean_control'),(4,'joint_feedback')]:
            data[row][1:]=[f'{v:.2f}' for v in tm['main_unrounded_values'][source].values()]
        data.append(['Relative improvement']+list(tm['relative_improvement']['display_percent'].values()))
    elif key in ['loss','published_benchmark']:pass
    else:data=[[format_num(v) for v in r] for r in data]
    if key=='sequences':data[0][3]='PA-MPJPE'
    wide=not supp
    if wide:section(doc,1)
    p=para(doc,f'Table {"S" if supp else ""}{num}. '+b['caption'],'Caption');p.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=0,cols=len(data[0]));t.autofit=False
    widths=([2.55]+[(FULL-2.55)/4]*4) if key=='main' else ([1.4,2.0,.55,.7,1.05,1.175] if key=='published_benchmark' else [FULL/len(data[0])]*len(data[0]))
    for i,row in enumerate(data):
        cells=t.add_row().cells
        for j,val in enumerate(row):
            cells[j].width=Inches(widths[j]);p=cells[j].paragraphs[0];p.paragraph_format.first_line_indent=Inches(0)
            p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(1 if supp else 2);p.paragraph_format.line_spacing=Pt(10)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or key=='published_benchmark' and j==1 else WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(cite(val));r.font.size=Pt(9)
            if i==0 or key=='main' and i==4:r.bold=True
            if key=='published_benchmark' and 0<i<len(data)-1 and j in [4,5]:
                r.bold=float(val)==min(float(v[j]) for v in data[1:-1])
            if key=='published_benchmark' and i==len(data)-1 and j==0:r.bold=True
            r.font.color.rgb=RGBColor(0,0,0)
            if key=='main' and i==5:r.font.color.rgb=RGBColor.from_string('B00000')
            tcPr=cells[j]._tc.get_or_add_tcPr();borders=el('w:tcBorders')
            tcPr.append(el('w:shd',val='clear',color='auto',fill='FFFFFF'))
            for edge in ['left','right','top','bottom']:borders.append(el('w:'+edge,val='nil'))
            if i==0:
                borders.append(el('w:top',val='single',sz='8',color='000000'));borders.append(el('w:bottom',val='single',sz='4',color='000000'))
            if i==len(data)-1:borders.append(el('w:bottom',val='single',sz='8',color='000000'))
            if key=='main' and i==4:borders.append(el('w:top',val='single',sz='4',color='000000'))
            if key=='published_benchmark' and i==len(data)-1:borders.append(el('w:top',val='single',sz='4',color='000000'))
            tcPr.append(borders)
        trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(el('w:cantSplit'))
        if i==0:trpr.append(el('w:tblHeader'))
    for i,row in enumerate(t.rows):
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.keep_with_next=i==0
    gap=doc.add_paragraph();gap.paragraph_format.space_before=gap.paragraph_format.space_after=Pt(0)
    gap.paragraph_format.line_spacing=Pt(1);gap.add_run().font.size=Pt(1)
    if wide:section(doc,2)

def ref_text(r):
    f=r['fields'];a=f['author'].replace("Gu\\'enol\\'e",'Guénolé').replace("L\\'aszl\\'o",'László')
    authors=a.split('; ') if ';' in a else a.split(' and ')
    authors=[' '.join(reversed(x.split(', ',1))) if ', ' in x else x for x in authors]
    a=(', '.join(authors[:-1])+', and '+authors[-1] if len(authors)>1 else authors[0]).rstrip('.')
    venue=f.get('booktitle',f.get('journal',''))
    if f.get('eprint'):venue='arXiv:'+f['eprint']
    tail=venue
    if f.get('volume'):tail+=' '+f['volume']+(f'({f["number"]})' if f.get('number') else '')
    if f.get('pages'):tail+=', pages '+f['pages'].replace('-', '–')
    return f'[{r["number"]}] {a}. {f["title"]}. {tail}, {f["year"]}.'

def build(name,supp=False):
    m=json.loads((H/('supplement_revised.json' if supp else 'manuscript_revised.json')).read_text('utf8'));d=Document();setup(d,supp)
    title=m.get('title','STaR-Body Supplementary Material');d.core_properties.title=title
    display_title=title.replace(' for Spatiotemporal','\nfor Spatiotemporal') if not supp else title
    p=para(d,display_title,'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p._p.get_or_add_pPr().append(el('w:suppressAutoHyphens'))
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(14);r.bold=True
    p=para(d,'Anonymous submission');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_after=Pt(12)
    if not supp:
        section(d,2);p=para(d,'Abstract');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].bold=True;p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_after=Pt(6)
        p=para(d,m['abstract']);p.paragraph_format.first_line_indent=Inches(0)
        for r in p.runs:r.italic=True
    fig=tab=0
    blocks=copy.deepcopy(m['blocks'])
    if not supp:
        overview=next(b for b in blocks if b.get('key')=='overview');blocks.remove(overview)
        idx=next(i for i,b in enumerate(blocks) if b.get('text')=='2. Related Work');blocks.insert(idx,overview)
        identity=next(b for b in blocks if b.get('key')=='identity');blocks.remove(identity)
        idx=next(i for i,b in enumerate(blocks) if b.get('text')=='3.2. Track-conditioned shape');blocks.insert(idx,identity)
        benchmark=next(b for b in blocks if b.get('key')=='published_benchmark');blocks.remove(benchmark)
        idx=next(i for i,b in enumerate(blocks) if b.get('text')=='4. Experiments');blocks.insert(idx,benchmark)
    for bi,b in enumerate(blocks):
        typ=b['type']
        if typ=='heading':d.add_heading(b['text'],level=b['level'])
        elif typ=='paragraph':para(d,b['text'],boldlead=b.get('lead_bold',False) or b['text'].split('. ')[0] in ['Image-based estimation and geometric feedback','Temporal and world-grounded reconstruction','Pose representations and observation conditioning','Visual representations'])
        elif typ=='contributions':
            for item in b['items']:
                p=para(d,'• '+item);p.paragraph_format.left_indent=Inches(.1);p.paragraph_format.first_line_indent=Inches(-.1);p.paragraph_format.space_after=Pt(2)
        elif typ=='equation':eq(d,b['tex'])
        elif typ=='figure':
            fig+=1
            adjacent_before=bi>0 and blocks[bi-1]['type']=='figure'
            adjacent_after=bi+1<len(blocks) and blocks[bi+1]['type']=='figure'
            figure(d,b,fig,supp,not adjacent_before,not adjacent_after)
        elif typ=='table':tab+=1;table(d,b,tab,supp)
    if not supp:
        d.add_heading('LLM Usage Statement',1)
        para(d,'OpenAI Codex assisted with manuscript editing, source and result cross-checking, reference metadata verification, and document preparation. Google Gemini supplied the overview reference; OpenAI image generation revised the overview and produced the landmark-feedback and frame-ownership drafts. The figures were subsequently edited and checked against the implemented computation. Human, mesh, and attention illustrations are schematic. Quantitative plots and tables use archived measurements and cited original-paper results; generative tools did not produce experimental outcomes.')
        d.add_heading('References',1)
        for r in refs:
            p=para(d,ref_text(r));p.paragraph_format.left_indent=Inches(.18);p.paragraph_format.first_line_indent=Inches(-.18);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=Pt(10)
            p.paragraph_format.keep_together=True
            for run in p.runs:run.font.size=Pt(9)
    dest=OUT/(name+'.docx');d.save(dest);print(dest)
if __name__ == '__main__':
    target = sys.argv[1] if len(sys.argv)>1 else 'all'
    if target in ['all','main']:build('STaR-Body_CVPR_2027')
    if target in ['all','supplement']:build('STaR-Body_Supplement',True)
